"""Render each deck in slides_data.py as a branded 16:9 PowerPoint file with speaker notes.

Run with: uv run --with python-pptx python make_pptx.py
"""

import pathlib
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from slides_data import DECKS, SLIDES, Slide

HERE = pathlib.Path(__file__).parent
OUT = HERE.parent / "downloads"

NAVY = RGBColor(0x05, 0x1E, 0x39)
NAVY2 = RGBColor(0x00, 0x3D, 0x69)
GOLD = RGBColor(0xB3, 0x90, 0x51)
GOLD_DARK = RGBColor(0x8F, 0x71, 0x3D)
GOLD_LIGHT = RGBColor(0xDE, 0xBD, 0x88)
DIPLOMA = RGBColor(0xF9, 0xF6, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0x3A, 0x4A, 0x5C)
MIST = RGBColor(0xD4, 0xD9, 0xDF)
TONES = {
    "Red": RGBColor(0xB3, 0x26, 0x1E),
    "Amber": RGBColor(0x9A, 0x5B, 0x00),
    "Green": RGBColor(0x06, 0x60, 0x34),
}
SLAB, SANS, MONO = "Roboto Slab", "Roboto", "Roboto Mono"
W, H = Inches(13.333), Inches(7.5)
LEFT, RIGHT = Inches(0.93), Inches(0.93)
WIDTH = W - LEFT - RIGHT


def fill(shape, color: RGBColor) -> None:
    """Give a shape a solid fill and no outline."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def text(
    slide,
    x,
    y,
    w,
    h,
    runs,
    size: int,
    color: RGBColor,
    font: str = SANS,
    bold: bool = False,
    anchor=MSO_ANCHOR.TOP,
):
    """Add a text box; ``runs`` is a list of paragraphs, each a str or a list of (text, color) runs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.45)
        for t, c in [(para, color)] if isinstance(para, str) else para:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.name, r.font.bold = Pt(size), font, bold
            r.font.color.rgb = c
    return tb


def gold_bar(slide) -> None:
    """Gold rule along the bottom edge, as on the GT web pages."""
    fill(
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, H - Inches(0.14), W, Inches(0.14)
        ),
        GOLD,
    )


def footer(slide, label: str, n: int, total: int) -> None:
    """Logo, deck label, and slide number along the bottom."""
    y = H - Inches(0.55)
    slide.shapes.add_picture(
        str(HERE / "assets" / "logo-navy.png"), LEFT, y, height=Inches(0.28)
    )
    text(
        slide,
        LEFT + Inches(1.9),
        y,
        Inches(7),
        Inches(0.3),
        [f"Applied AI at TIYA Services · {label}"],
        11,
        SOFT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        slide,
        W - RIGHT - Inches(1.5),
        y,
        Inches(1.5),
        Inches(0.3),
        [f"{n} / {total}"],
        11,
        SOFT,
        MONO,
        anchor=MSO_ANCHOR.MIDDLE,
    ).text_frame.paragraphs[0].alignment = 3


def tag(slide, label: str, color: RGBColor) -> None:
    """Pill label above the title (Exercise · 15 min, Live demo)."""
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        LEFT,
        Inches(0.7),
        Inches(0.25) + Pt(10.5 * 0.72 * len(label)),
        Inches(0.38),
    )
    fill(pill, color)
    tf = pill.text_frame
    tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run()
    r.text = label.upper()
    r.font.size, r.font.bold, r.font.name = Pt(12), True, SANS
    r.font.color.rgb = WHITE


def render(prs, s: Slide, n: int, total: int, label: str) -> None:
    """Add one slide to the presentation."""
    ty, title, bullets, note = s
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background.fill
    if ty == "title":
        bg.solid()
        bg.fore_color.rgb = NAVY
        sl.shapes.add_picture(
            str(HERE / "assets" / "logo-white.png"),
            LEFT,
            Inches(0.6),
            height=Inches(0.55),
        )
        text(
            sl,
            LEFT,
            Inches(1.85),
            WIDTH,
            Inches(0.4),
            [f"APPLIED AI AT TIYA SERVICES · {label.upper()}"],
            15,
            GOLD_LIGHT,
            SANS,
        )
        text(sl, LEFT, Inches(2.35), WIDTH, Inches(1.8), [title], 54, WHITE, SLAB, True)
        text(sl, LEFT, Inches(5.0), WIDTH, Inches(1.2), bullets, 20, MIST)
        gold_bar(sl)
    elif ty == "section":
        bg.gradient()
        bg.gradient_angle = 0
        bg.gradient_stops[0].color.rgb = NAVY
        bg.gradient_stops[1].color.rgb = NAVY2
        name, _, when = title.partition(" · ")
        text(
            sl,
            LEFT,
            Inches(2.5),
            WIDTH,
            Inches(0.4),
            [(when or label).upper()],
            15,
            GOLD_LIGHT,
            SANS,
        )
        text(sl, LEFT, Inches(3.0), WIDTH, Inches(1.6), [name], 48, WHITE, SLAB, True)
        text(sl, LEFT, Inches(4.7), WIDTH, Inches(0.8), bullets, 20, MIST)
        gold_bar(sl)
    else:
        bg.solid()
        bg.fore_color.rgb = DIPLOMA if ty == "exercise" else WHITE
        top = Inches(0.75)
        if ty == "exercise":
            m = re.match(r"(\d+) min", note)
            tag(sl, f"Exercise · {m.group(0)}" if m else "Exercise", GOLD_DARK)
            top = Inches(1.35)
        elif ty == "demo":
            tag(sl, "Live demo", NAVY)
            top = Inches(1.35)
        text(sl, LEFT, top, WIDTH, Inches(1.0), [title], 36, NAVY, SLAB, True)
        body_top = top + Inches(1.15)
        if ty == "columns":
            n_cols, gap = len(bullets), Inches(0.3)
            cw = Emu(int((WIDTH - gap * (n_cols - 1)) / n_cols))
            small = n_cols >= 5
            for k, b in enumerate(bullets):
                head, _, body = b.partition(": ")
                tone = TONES.get(head, GOLD_DARK)
                x = LEFT + (cw + gap) * k
                fill(
                    sl.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE, x, body_top, cw, Inches(0.06)
                    ),
                    TONES.get(head, GOLD),
                )
                text(
                    sl,
                    x,
                    body_top + Inches(0.22),
                    cw,
                    Inches(0.8),
                    [head],
                    19 if small else 24,
                    tone,
                    SLAB,
                    True,
                )
                text(
                    sl,
                    x,
                    body_top + Inches(1.0),
                    cw,
                    Inches(2.6),
                    [body],
                    15 if small else 19,
                    NAVY,
                )
        else:
            paras = [[("•  ", GOLD), (b, NAVY)] for b in bullets]
            text(sl, LEFT, body_top, WIDTH, Inches(4.2), paras, 22, NAVY)
        footer(sl, label, n, total)
    if note:
        sl.notes_slide.notes_text_frame.text = note


def build(key: str) -> None:
    """Write downloads/<key>-slides.pptx."""
    meta, slides = DECKS[key], SLIDES[key]
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    label = f"{meta['label']} · {meta['title']}"
    for i, s in enumerate(slides):
        render(prs, s, i + 1, len(slides), label)
    OUT.mkdir(exist_ok=True)
    path = OUT / f"{key}-slides.pptx"
    prs.save(str(path))
    print(f"{path.name}  {len(slides)} slides  {path.stat().st_size // 1024} KB")


if __name__ == "__main__":
    for k in DECKS:
        build(k)
